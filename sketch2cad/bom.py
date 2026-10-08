"""Bill of materials from a drawing spec → Excel (improvement 5)."""
from __future__ import annotations

import math
from pathlib import Path

LINEAR = ("pipe", "trench_drain")
COLUMNS = (("no", "No", "מס'"), ("item", "Item", "פריט"), ("en", "Description", "תיאור באנגלית"),
           ("he", "תיאור", "תיאור"), ("qty", "Qty", "כמות"), ("unit", "Unit", "יחידה"),
           ("part_no", "Part no.", "מק\"ט"), ("supplier", "Supplier", "ספק"),
           ("status", "Status", "סטטוס"), ("source", "Source", "מקור"))


def _length(points) -> float:
    return sum(math.dist(a, b) for a, b in zip(points, points[1:]))


def _to_m(value, spec) -> float:
    return value if spec.get("units") == "m" else value / 1000


def build_bom(spec: dict, catalog=None) -> list[dict]:
    """One row per (catalog item, status); pipes and trench drains in metres, everything else in pieces."""
    legend_no = {l["item"]: l.get("no") for l in spec.get("legend", []) if l.get("item")}
    rows: dict[tuple, dict] = {}
    elements = list(spec.get("elements", []))
    elements += [{"type": "cover", "item": e["cover_item"], "status": e.get("cover_status", e.get("status", "new"))}
                 for e in elements if e.get("cover_item")]          # manhole covers are separate BOM lines
    for el in elements:
        item_id = el.get("item")
        if not item_id:
            continue
        item = (catalog.get(item_id) if catalog else None) or {}
        key = (item_id, el.get("status", "new"))
        row = rows.setdefault(key, {
            "no": legend_no.get(item_id, ""), "item": item_id,
            "en": (item.get("name") or {}).get("en", item_id), "he": (item.get("name") or {}).get("he", ""),
            "qty": 0.0, "unit": item.get("unit", "m" if el["type"] in LINEAR else "pcs"),
            "part_no": item.get("part_no", ""), "supplier": item.get("supplier", item.get("manufacturer", "")),
            "status": key[1], "source": (item.get("source") or {}).get("submittal", "")})
        if row["unit"] == "m" and el.get("points"):
            row["qty"] += _to_m(_length(el["points"]), spec)
        elif row["unit"] == "m" and el["type"] == "bend":
            row["qty"] += _to_m(math.radians(abs(el["end_angle"] - el["start_angle"])) * el["radius"], spec)
        else:
            row["qty"] += 1
    for l in spec.get("legend", []):            # legend-only items (e.g. bolts, gaskets)
        item_id = l.get("item")
        if item_id and not any(k[0] == item_id for k in rows):
            item = (catalog.get(item_id) if catalog else None) or {}
            rows[(item_id, "new")] = {
                "no": l.get("no", ""), "item": item_id, "en": l.get("en", ""), "he": l.get("he", ""),
                "qty": float(len(l.get("targets") or [1])), "unit": item.get("unit", "pcs"),
                "part_no": item.get("part_no", ""), "supplier": item.get("supplier", ""),
                "status": "new", "source": (item.get("source") or {}).get("submittal", "")}
    out = sorted(rows.values(), key=lambda r: (str(r["no"]).zfill(4) if r["no"] != "" else "zzzz", r["item"]))
    for r in out:
        r["qty"] = round(r["qty"], 2) if r["unit"] == "m" else int(round(r["qty"]))
    return out


def bom_to_xlsx(rows: list[dict], path, *, lang: str = "en", title: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.sheet_view.rightToLeft = lang == "he"
    start = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=13)
        start = 3
    head = 2 if lang == "he" else 1
    for c, col in enumerate(COLUMNS, 1):
        cell = ws.cell(start, c, col[head])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start + 1):
        for c, col in enumerate(COLUMNS, 1):
            ws.cell(r, c, row.get(col[0], ""))
    for c, width in enumerate((6, 24, 50, 40, 8, 7, 16, 16, 10, 16), 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = ws.cell(start + 1, 1)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
