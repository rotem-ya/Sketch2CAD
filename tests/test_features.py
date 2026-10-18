"""Catalog, checks, BOM, typical details, submittal package, print set and markup."""
import json

import pymupdf
import pytest

from sketch2cad import bom, checks, typical
from sketch2cad.catalog import Catalog, load_catalog, save_item
from sketch2cad.spec import ELEMENT_TYPES
from sketch2cad.documents import create_document
from sketch2cad.package import build_package
from sketch2cad.printset import build_print_set
from sketch2cad.projects import create_project


@pytest.fixture(autouse=True)
def isolated_home(tmp_path, monkeypatch):
    monkeypatch.setenv("SKETCH2CAD_HOME", str(tmp_path / "home"))


@pytest.fixture
def cat():
    return load_catalog()


def codes(findings):
    return [f.code for f in findings]


# ---------------------------------------------------------------- catalog
def test_catalog_seed_and_sources(cat):
    avk = cat.get("avk_0661_dn150")
    assert avk["dims"]["L"] == 210 and avk["flange"]["PCD"] == 240
    assert avk["name"]["he"].startswith("מגוף")
    assert Catalog.missing_source(avk) == ["t"]                       # t comes from the Plasson data
    assert Catalog.missing_source(cat.get("bolt_m20_ss316")) == ["source"]
    assert {i["id"] for i in cat.items("pipe")} >= {"pexgol_160_c10", "pipe_bt_400"}


def test_project_catalog_overrides(tmp_path):
    p = create_project(tmp_path, "20117", "חצרים", "Hatzerim")
    save_item({"id": "golan_pex5081_160", "category": "connector", "name": {"en": "Golan", "he": "גולן"},
               "dims": {"L": 240}, "source": {"submittal": "57 00 00-8.1", "fields": ["L"]}}, p)
    c = load_catalog(p)
    assert c.get("golan_pex5081_160")["dims"]["L"] == 240
    assert Catalog.missing_source(c.get("golan_pex5081_160")) == []


# ---------------------------------------------------------------- checks
def test_cover_check(cat):
    ok = typical.generate("riser_connection", catalog=cat)
    assert "cover" not in codes(checks.run_checks(ok, cat))
    shallow = typical.generate("riser_connection", {"cover": 900}, catalog=cat)
    found = [f for f in checks.run_checks(shallow, cat) if f.code == "cover"]
    assert found and found[0].severity == "error" and "900" in found[0].en


def test_flange_and_bolt_checks(cat):
    spec = typical.generate("riser_connection", catalog=cat)
    assert "flange_drilling" in codes(checks.run_checks(spec, cat))   # Golan ASA150 vs PN16
    spec["legend"] = [l for l in spec["legend"] if l.get("item") != "bolt_m20_ss316"]
    assert "bolts_missing" in codes(checks.run_checks(spec, cat))
    c = Catalog({i["id"]: i for i in cat.items()})
    c.add({"id": "bolt_cs", "category": "bolt", "material": "Carbon steel 8.8"})
    spec["legend"].append({"no": 9, "en": "BOLTS", "he": "ברגים", "item": "bolt_cs"})
    assert "bolts_material" in codes(checks.run_checks(spec, c))


def test_discharge_distance(cat):
    spec = typical.generate("trench_to_manhole", {"run_length": 10.0}, catalog=cat)
    assert "discharge_distance" not in codes(checks.run_checks(spec, cat))
    spec = typical.generate("trench_to_manhole", {"run_length": 30.0, "max_discharge_distance": 20.0}, catalog=cat)
    assert "discharge_distance" in codes(checks.run_checks(spec, cat))
    spec["elements"] = [e for e in spec["elements"] if e["type"] not in ("manhole", "pipe")]
    assert "no_outlet" in codes(checks.run_checks(spec, cat))


def test_unknown_item(cat):
    spec = {"units": "mm", "elements": [{"id": "x", "type": "pipe", "item": "nope", "points": [[0, 0], [1, 0]]}]}
    assert checks.run_checks(spec, cat)[0].code == "unknown_item"


# ---------------------------------------------------------------- BOM
def test_bom_quantities(cat, tmp_path):
    spec = typical.generate("valve_in_line", {"pipe_length": 1000}, catalog=cat)
    rows = {r["item"]: r for r in bom.build_bom(spec, cat)}
    assert rows["steel_pipe_6in_sch40"]["qty"] == 2.0 and rows["steel_pipe_6in_sch40"]["unit"] == "m"
    assert rows["flange_6in_pn16"]["qty"] == 2 and rows["avk_0661_dn150"]["qty"] == 1
    assert rows["avk_0661_dn150"]["source"] == "07 00 00-7.1"
    out = bom.bom_to_xlsx(list(rows.values()), tmp_path / "bom.xlsx", lang="he", title="20117")
    from openpyxl import load_workbook
    ws = load_workbook(out).active
    assert ws.sheet_view.rightToLeft and ws.cell(3, 1).value == "מס'"


def test_bom_metres_for_m_units(cat):
    spec = typical.generate("trench_to_manhole", {"run_length": 10.0}, catalog=cat)
    rows = {r["item"]: r for r in bom.build_bom(spec, cat)}
    assert rows["trench_drain_generic"]["unit"] == "m" and 9 < rows["trench_drain_generic"]["qty"] < 10


# ---------------------------------------------------------------- typicals
def test_typicals_generate_valid_specs(cat):
    names = [t.name for t in typical.list_typicals()]
    assert names == ["riser_connection", "trench_to_manhole", "valve_in_line"]
    for name in names:
        spec = typical.generate(name, catalog=cat)
        types = {e["type"] for e in spec["elements"]}
        assert types <= set(ELEMENT_TYPES)
        assert all(l["he"] and l["en"] for l in spec["legend"])
        assert len({e["id"] for e in spec["elements"]}) == len(spec["elements"])


# ---------------------------------------------------------------- package + print set
@pytest.fixture
def project_with_doc(tmp_path, cat):
    p = create_project(tmp_path / "projects", "20117", "חצרים", "Hatzerim", client="USACE")
    d = create_document(p, title_he="חיבור לגמל קיים", title_en="Connection to existing camel",
                        discipline="W", subject="CAMEL")
    spec = typical.generate("riser_connection", catalog=cat)
    (d.path / "spec.json").write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    drawing = pymupdf.open()
    drawing.new_page(width=1191, height=842).insert_text((72, 72), "DRAWING SHEET")
    drawing.save(d.path / f"{d.code}.pdf")
    d.outputs["pdf"] = f"{d.code}.pdf"
    d.save()
    sheet = pymupdf.open()
    sheet.new_page().insert_text((72, 72), "AVK DATASHEET")
    sheet.save(p.path / "00_Sources" / "0661_standard_Israel.pdf")
    return p, d


@pytest.mark.parametrize("lang", ["en", "he"])
def test_package(project_with_doc, cat, tmp_path, lang):
    p, d = project_with_doc
    out = build_package(p, d, cat, tmp_path / f"pkg_{lang}.pdf", lang=lang)
    pdf = pymupdf.open(out)
    text = "".join(page.get_text() for page in pdf)
    assert "DRAWING SHEET" in text and "AVK DATASHEET" in text
    assert "07 00 00-7.1" in text and "20117-W-CAMEL-01" in text
    assert ("SUBMITTAL PACKAGE" if lang == "en" else "הגשה") in text


def test_print_set(project_with_doc, tmp_path):
    p, d = project_with_doc
    create_document(p, title_en="No pdf yet", title_he="ללא", discipline="W", subject="X")
    out = build_print_set(p, tmp_path / "set.pdf", lang="en")
    pdf = pymupdf.open(out)
    text = pdf[0].get_text()
    assert "20117-W-CAMEL-01" in text and "no PDF" in text
    assert "DRAWING SHEET" in pdf[1].get_text()
    assert [t[1].split(" ")[0] for t in pdf.get_toc()] == ["Sheet", "20117-W-CAMEL-01"]


def test_pick_folder_without_desktop(monkeypatch):
    from sketch2cad import native
    monkeypatch.delenv("DISPLAY", raising=False)
    assert native.pick_folder() is None


# ---------------------------------------------------------------- self-update
def _zip_of(files: dict) -> bytes:
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, text in files.items():
            z.writestr(f"Sketch2CAD-abc123/{name}", text)
    return buf.getvalue()


def test_updater_latest_and_apply(tmp_path, monkeypatch):
    from sketch2cad import updater
    api = {"sha": "abc123", "commit": {"committer": {"date": "2026-10-08T10:00:00Z"},
                                       "message": "New feature\n\nbody"}}
    monkeypatch.setattr(updater, "_get", lambda url, timeout: json.dumps(api).encode())
    rel = updater.latest()
    assert (rel.sha, rel.date, rel.message) == ("abc123", "2026-10-08", "New feature")

    app = tmp_path / "app"
    (app / ".venv").mkdir(parents=True)
    (app / ".venv" / "keep.txt").write_text("venv")
    (app / "requirements.txt").write_text("ezdxf\n")
    (app / "old.py").write_text("old")
    assert updater.installed_version(app) is None and updater.update_available(rel, app)
    calls = []
    monkeypatch.setattr(updater.subprocess, "run", lambda cmd, **kw: calls.append(cmd))
    zip_bytes = _zip_of({"app/app.py": "new app", "requirements.txt": "ezdxf\n", ".venv/keep.txt": "NO"})
    sha = updater.apply_update(app, "abc123", log=lambda m: None, download=lambda url, t: zip_bytes)
    assert sha == "abc123" and (app / "app" / "app.py").read_text() == "new app"
    assert (app / ".venv" / "keep.txt").read_text() == "venv"            # never touched
    assert calls == []                                                     # requirements unchanged → no pip
    assert updater.installed_version(app) == "abc123" and not updater.update_available(rel, app)
    zip_bytes = _zip_of({"requirements.txt": "ezdxf\nnewlib\n"})
    updater.apply_update(app, "def456", log=lambda m: None, download=lambda url, t: zip_bytes)
    assert calls and calls[0][1:4] == ["-m", "pip", "install"]



# ---------------------------------------------------------------- markup + revise on request (no network)
def _white_png(path, w=400, h=300):
    from PIL import Image
    Image.new("RGB", (w, h), "white").save(path)
    return path


def test_markup_clean_and_compose(tmp_path):
    from PIL import Image
    from sketch2cad import markup
    strokes = [
        {"tool": "line", "color": "red", "width": 0.01, "points": [[0.1, 0.5], [0.3, 0.5], [0.9, 0.5]]},
        {"tool": "pen", "color": "blue", "points": [[0.5, 0.1], [0.5, 0.2], [0.5, 0.3]]},
        {"tool": "laser", "points": [[0, 0], [1, 1]]},                         # unknown tool: dropped
        {"tool": "pen", "color": "pink", "points": [[2, -1], [0.2, 0.2]]},     # colour -> red, clipped
        {"tool": "pen", "points": [[0.5, 0.5]]},                                # a single point: dropped
    ]
    clean = markup.clean_strokes(strokes)
    assert [s["tool"] for s in clean] == ["line", "pen", "pen"]
    assert clean[0]["points"] == [[0.1, 0.5], [0.9, 0.5]]                      # a line keeps its two ends
    assert clean[2]["color"] == "red" and clean[2]["points"][0] == [1.0, 0.0]
    out = markup.compose(_white_png(tmp_path / "p.png"), strokes, tmp_path / "m" / "marked.png")
    with Image.open(out) as img:
        assert img.getpixel((200, 150)) == markup.COLORS["red"]               # on the red line
        assert img.getpixel((200, 45)) == markup.COLORS["blue"]               # on the blue stroke
        assert img.getpixel((20, 280)) == (255, 255, 255)                      # untouched
    assert markup.canvas_image(tmp_path / "p.png").startswith("data:image/jpeg;base64,")

